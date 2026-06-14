"""
Generates docs/scoring-refactor/WiFi_Risk_Score_Methodology.docx

One-shot documentation generator for the Why-PII? WiFi Risk Score
methodology, written for an expert-review audience (not developers).
Re-run after editing this script to regenerate the .docx.
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = "WiFi_Risk_Score_Methodology.docx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
GREY = RGBColor(0x59, 0x59, 0x59)


def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY
    return h


def add_caption(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GREY
    return p


def add_table(doc, headers, rows, col_widths=None, header_bg="1F2A44"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(hdr_cells[i], header_bg)

    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    return table


def main():
    doc = Document()

    # ---- base font ----
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ---- title page ----
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Why-PII?")
    run.bold = True
    run.font.size = Pt(34)
    run.font.color.rgb = NAVY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("WiFi Risk Score Methodology")
    run.font.size = Pt(20)
    run.font.color.rgb = GREY

    subtitle2 = doc.add_paragraph()
    subtitle2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle2.add_run("Formula, Scoring Algorithm, and Standards Basis")
    run.font.size = Pt(13)
    run.italic = True
    run.font.color.rgb = GREY

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "Document version 1.0 — June 2026\n"
        "Prepared for expert review of the Wi-Fi network risk scoring methodology"
    )
    meta_run.font.size = Pt(10)
    meta_run.font.color.rgb = GREY

    doc.add_page_break()

    # ---- 1. purpose ----
    add_heading(doc, "1. Purpose of This Document", level=1)
    doc.add_paragraph(
        "Why-PII? is a Wi-Fi security assessment tool. For every network it scans, "
        "the system produces a single, easy-to-read “overall risk score” on a 0–100 "
        "scale, together with a risk category — LOW, MEDIUM, HIGH, or CRITICAL. This "
        "score is shown to end users on the captive portal and to administrators on "
        "the dashboard, and it drives recommendations and portal warnings."
    )
    doc.add_paragraph(
        "This document explains, for an expert (non-developer) audience, exactly how "
        "that score is calculated: what inputs feed into it, the mathematical formula "
        "used to combine those inputs, why that formula was chosen over the "
        "alternatives, and which published standards and frameworks support the "
        "approach. It is intended to support an independent evaluation of whether the "
        "scoring methodology is sound, defensible, and fit for purpose."
    )

    # ---- 2. inputs ----
    add_heading(doc, "2. What Feeds Into the Score", level=1)
    doc.add_paragraph(
        "Every scan checks a Wi-Fi access point (AP) against a fixed WFVT Data Set of "
        "seven known Wi-Fi vulnerabilities and threats (“WFVT” items). Each item has "
        "two properties that matter for scoring:"
    )

    add_table(
        doc,
        headers=["ID", "Item", "Risk Score (CVSS v4.0 Base)"],
        rows=[
            ["WFVT-001", "Open System Authentication", "9.4"],
            ["WFVT-002", "Weak or Deprecated Cryptographic Algorithms", "9.4"],
            ["WFVT-003", "Wi-Fi Protected Setup (WPS) Enabled", "7.6"],
            ["WFVT-004", "Protected Management Frames (PMF) Disabled", "7.1"],
            ["WFVT-005", "Evil Twin Attack", "9.3"],
            ["WFVT-006", "Deauthentication Attack", "8.5"],
            ["WFVT-007", "MAC Address Spoofing Attack", "9.4"],
            ["", "Maximum Possible Risk (Rmax)", "60.7"],
        ],
        col_widths=[1.1, 4.0, 1.5],
    )
    doc.add_paragraph()
    add_caption(
        doc,
        "Table 1. The fixed WFVT Data Set with CVSS v4.0 Base Score severity "
        "values used throughout this document."
    )
    doc.add_paragraph()

    p = doc.add_paragraph(style="List Bullet")
    p.add_run("Presence (Pi).").bold = True
    p.add_run(
        " Whether the item is actually detected on this specific network at this "
        "point in time. This is a yes/no determination, not a guess — it comes "
        "directly from the scan and from live, continuous detection while the "
        "session runs."
    )

    p = doc.add_paragraph(style="List Bullet")
    p.add_run("Severity (CVSSi).").bold = True
    p.add_run(
        " The intrinsic severity of that item, expressed as a CVSS (Common "
        "Vulnerability Scoring System) Base Score from 0.0 to 10.0. These scores are "
        "fixed reference values taken from the CVSS v4.0 vector for each "
        "vulnerability/threat type and do not change from scan to scan."
    )

    doc.add_paragraph(
        "The seven WFVT Data Set items (WFVT-001 through WFVT-007) cover the categories "
        "most commonly flagged in Wi-Fi security guidance: open/weak authentication, "
        "weak encryption, WPS exposure, disabled management-frame protection, and "
        "active attack techniques such as evil-twin access points, deauthentication "
        "attacks, and MAC spoofing."
    )

    doc.add_paragraph(
        "Two categories of items are evaluated slightly differently, reflecting how "
        "they are detected:"
    )
    p = doc.add_paragraph(style="List Bullet")
    p.add_run("Configuration vulnerabilities").bold = True
    p.add_run(
        " (e.g. open authentication, WPS enabled) are present if the scan finds the "
        "corresponding misconfiguration on the access point — these are static "
        "properties of how the network is configured."
    )
    p = doc.add_paragraph(style="List Bullet")
    p.add_run("Active threats").bold = True
    p.add_run(
        " (e.g. evil-twin AP, deauthentication attack, MAC spoofing) are present "
        "only if continuous monitoring has most recently classified them as ACTIVE. "
        "If a threat is later cleared, it stops contributing to the score on the "
        "next recalculation — the score is always a live snapshot, not a "
        "ratchet that only goes up."
    )

    # ---- 3. old formula ----
    add_heading(doc, "3. Previous Formula and Its Limitation", level=1)
    doc.add_paragraph(
        "The original scoring formula combined the presence and severity of each "
        "WFVT Data Set item using a simple weighted sum, normalised against the maximum "
        "possible total:"
    )

    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Score % = ( Σ Pi × CVSSi ) / Rmax × 100")
    r.bold = True
    r.font.size = Pt(13)

    doc.add_paragraph(
        "Here, Rmax = 60.7 is the fixed sum of the CVSS Base Scores of all seven "
        "WFVT Data Set items — i.e. the theoretical “worst case” total if every "
        "single item were present at once."
    )

    doc.add_paragraph(
        "The problem with this approach is that the fixed denominator dilutes the "
        "impact of any single severe finding. For example, a network with only "
        "“Open Authentication” detected — a vulnerability independently rated as "
        "Critical (CVSS 9.4 of 10) — produced:"
    )
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("9.4 / 60.7 × 100 ≈ 15.5%  →  LOW")
    r.bold = True

    doc.add_paragraph(
        "A user viewing “15% / Low” on the captive portal would reasonably assume "
        "the network is relatively safe — despite the presence of a vulnerability "
        "that, on its own, is rated Critical by the CVSS standard. This mismatch "
        "between the underlying severity of a single finding and the headline score "
        "was the primary motivation for revisiting the formula."
    )

    # ---- 4. new formula ----
    add_heading(doc, "4. New Formula: Noisy-OR / Probabilistic Combination", level=1)
    doc.add_paragraph(
        "The revised formula treats each detected item's severity as a probability "
        "of compromise, and combines all currently-present items using the standard "
        "formula for the probability that at least one of several independent "
        "events occurs:"
    )

    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Score % = ( 1 − Π (1 − CVSSi / 10) ) × 100")
    r.bold = True
    r.font.size = Pt(13)
    fp2 = doc.add_paragraph()
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = fp2.add_run("for every WFVT Data Set item i that is currently present (Pi = 1)")
    r2.italic = True
    r2.font.size = Pt(10)
    r2.font.color.rgb = GREY

    doc.add_paragraph(
        "Reading this formula in plain terms: each finding's CVSS score, divided by "
        "10, is treated as the likelihood that this finding alone is severe enough "
        "to lead to a compromise. The term (1 − CVSSi/10) is then the likelihood "
        "that this one finding does not lead to compromise. Multiplying these "
        "values together (Π, “product of”) gives the likelihood that none of "
        "the present findings lead to compromise. Subtracting that from 1 gives the "
        "likelihood that at least one of them does — which becomes the headline "
        "risk score."
    )

    add_heading(doc, "4.1 Why This Formula Behaves Better", level=2)
    for bullet in [
        "It is naturally bounded between 0% and just under 100% — no arbitrary "
        "scaling constant (like Rmax = 60.7) is required.",
        "It increases monotonically as more findings become present: adding a new "
        "active finding can only raise the score, never lower it (for a fixed set "
        "of other findings).",
        "A single very severe finding immediately produces a high score on its "
        "own — it is not “diluted” by being divided across a larger "
        "denominator.",
        "Multiple findings compound towards 100% rather than plateauing at the "
        "value of the single worst finding — a network with four serious "
        "issues is scored meaningfully higher than one with only the worst of "
        "those four.",
        "If no WFVT Data Set items are present, the product over an empty set is 1 by "
        "definition, giving a score of exactly 0% — a clean “no findings, no "
        "risk” baseline.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")

    # ---- 5. worked examples ----
    add_heading(doc, "5. Worked Examples", level=1)
    doc.add_paragraph(
        "The table below compares the previous and new formulas across four "
        "scenarios, using representative CVSS Base Scores for Open Authentication "
        "(9.4), Evil Twin AP (9.3), Deauthentication Attack (8.5), and Protected "
        "Management Frames Disabled (7.1)."
    )

    add_table(
        doc,
        headers=["Detected findings", "Previous formula (÷ 60.7)", "New formula (noisy-OR)"],
        rows=[
            ["None", "0%  →  LOW", "0%  →  LOW"],
            ["Open Authentication (9.4) only", "15.5%  →  LOW", "94%  →  CRITICAL"],
            [
                "Open Auth + Evil Twin + Deauth\n(9.4, 9.3, 8.5)",
                "44.8%  →  MEDIUM",
                "≈ 99.9%  →  CRITICAL",
            ],
            [
                "Open Auth + Evil Twin + Deauth\n+ PMF Disabled (9.4, 9.3, 8.5, 7.1)",
                "56.5%  →  MEDIUM",
                "≈ 99.98%  →  CRITICAL",
            ],
        ],
        col_widths=[2.6, 1.9, 1.9],
    )
    doc.add_paragraph()

    doc.add_paragraph(
        "Row 2 illustrates the central correction: a network with only one "
        "Critical-rated vulnerability (Open Authentication, CVSS 9.4) is no longer "
        "reported as “Low” risk. Rows 3 and 4 show that, unlike a simpler "
        "“use the worst CVSS score only” alternative — which would assign all "
        "three of rows 2–4 an identical 94% — the noisy-OR formula still "
        "distinguishes between a network with one severe issue and a network with "
        "several, by pushing the combined score closer to the 90–100 ceiling as "
        "more findings accumulate."
    )

    # ---- 6. real scan comparison ----
    add_heading(doc, "6. Live Verification Against Real Scans", level=1)
    doc.add_paragraph(
        "The new formula has been deployed and verified against real scan data "
        "captured by the field probe. The table below shows the same access point "
        "(SSID “Kerbs pisowifi”, BSSID 20:0D:B0:7A:BB:CC), with identical detected "
        "findings, scored once under the previous formula and once under the new "
        "formula:"
    )

    add_table(
        doc,
        headers=["Scan", "Findings detected", "Score (previous)", "Score (new)"],
        rows=[
            [
                "Mar 9, 2026, 05:33 AM",
                "Open Authentication (9.4, CRITICAL)\nWPS Enabled (7.6, HIGH)",
                "28  →  LOW",
                "—",
            ],
            [
                "Jun 14, 2026, 09:26 PM",
                "Open Authentication (9.4, CRITICAL)\nWPS Enabled (7.6, HIGH)",
                "—",
                "99  →  CRITICAL",
            ],
        ],
        col_widths=[1.6, 2.6, 1.1, 1.1],
    )
    doc.add_paragraph()

    doc.add_paragraph("Verification of the new score:")
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(
        "1 − (1 − 0.94) × (1 − 0.76) = 1 − (0.06 × 0.24) "
        "= 1 − 0.0144 = 0.9856  →  99% → CRITICAL"
    )
    r.italic = True

    doc.add_paragraph(
        "This pair of scans is the same access point with the same two detected "
        "vulnerabilities — only the scoring formula differs between the two "
        "scan dates, isolating the effect of the formula change from any change in "
        "the network itself."
    )

    # ---- 7. severity categories ----
    add_heading(doc, "7. Translating the Score Into a Risk Category", level=1)
    doc.add_paragraph(
        "The 0–100 score produced by either formula is mapped to one of four "
        "risk categories shown to users. This mapping is unchanged by the formula "
        "update — only the underlying score calculation changed, not the "
        "category boundaries:"
    )

    add_table(
        doc,
        headers=["Score range", "Risk category", "Typical meaning"],
        rows=[
            ["0", "LOW (None)", "No WFVT Data Set findings detected"],
            ["1 – 39", "LOW", "Minor or low-severity findings only"],
            ["40 – 69", "MEDIUM", "Moderate exposure; attention recommended"],
            ["70 – 89", "HIGH", "Significant exposure; remediation advised"],
            ["90 – 100", "CRITICAL", "Severe exposure; immediate action advised"],
        ],
        col_widths=[1.3, 1.5, 3.6],
    )
    doc.add_paragraph()

    # ---- 8. standards basis ----
    add_heading(doc, "8. Standards and Published Basis", level=1)
    doc.add_paragraph(
        "The methodology draws on four established sources. Each is summarised "
        "below together with how it specifically supports the design of this "
        "formula."
    )

    add_heading(doc, "8.1 CVSS v4.0 (FIRST.org — Forum of Incident Response and "
                      "Security Teams)", level=2)
    doc.add_paragraph(
        "The Common Vulnerability Scoring System (CVSS) v4.0 specification states "
        "that “the Base Score reflects the severity of a vulnerability according "
        "to its intrinsic characteristics which are constant over time.”"
    )
    doc.add_paragraph(
        "The per-item CVSSi values used in this methodology are exactly these "
        "intrinsic CVSS v4.0 Base Scores — they are not modified or re-weighted. "
        "The formula change affects only how multiple already-scored findings for "
        "a single network are combined into one overall number. Notably, the CVSS "
        "specification deliberately does not define a method for aggregating "
        "scores across multiple findings on a single asset — it explicitly leaves "
        "that decision to the implementing organisation. Choosing a different "
        "combination function is therefore within the intended scope of CVSS, "
        "whereas altering the individual CVSSi values themselves would not be."
    )

    add_heading(doc, "8.2 NIST SP 800-30 Rev. 1 — Guide for Conducting Risk "
                      "Assessments", level=2)
    doc.add_paragraph(
        "NIST Special Publication 800-30 Rev. 1 provides guidance on aggregating "
        "risk from multiple threat/vulnerability pairs affecting the same asset. "
        "Its guidance holds that such aggregated risk is not simply additive — "
        "combined risk from multiple contributing factors can exceed what any one "
        "factor would suggest on its own, reflecting compounding exposure."
    )
    doc.add_paragraph(
        "This directly supports moving away from a linear, fixed-denominator "
        "normalisation (which, as shown in Section 3, can understate the "
        "significance of a single severe finding) towards a combination function "
        "that compounds multiple findings together. The noisy-OR formula has "
        "exactly this property: it grows towards 100% as more high-severity "
        "findings accumulate, and — critically — never produces a combined "
        "score lower than the contribution of the single worst finding on its own. "
        "The presence-based likelihood model (a finding either is or is not "
        "currently observed, based on scan/detection evidence) is also consistent "
        "with NIST's evidence-based approach to assigning likelihood."
    )

    add_heading(doc, "8.3 OWASP Risk Rating Methodology", level=2)
    doc.add_paragraph(
        "OWASP's risk rating methodology expresses risk as the product of "
        "likelihood and impact (Risk = Likelihood × Impact). The "
        "presence-times-severity relationship for each individual finding in this "
        "methodology mirrors that structure directly: presence (Pi) represents "
        "likelihood, and the CVSS Base Score (CVSSi) represents impact/severity."
    )
    doc.add_paragraph(
        "OWASP's threat-modelling guidance further emphasises that an attacker "
        "typically needs only one successful exploit path to compromise a system. "
        "The noisy-OR formula expresses exactly this idea at the network level: "
        "“what is the probability that at least one of these N findings is "
        "successfully exploited?” — rather than treating findings as if their "
        "risks must all contribute additively and in proportion to a fixed total."
    )

    add_heading(doc, "8.4 Probability Theory and Reliability Engineering "
                      "(Fault Tree Analysis)", level=2)
    doc.add_paragraph(
        "The expression 1 − Π(1 − pi) is the standard formula from "
        "probability theory for the probability of the union of several "
        "independent events — i.e. the probability that at least one of several "
        "independent things happens. The same expression is used as the “OR-gate” "
        "combination rule in Fault Tree Analysis (FTA), a long-established "
        "reliability-engineering technique for combining multiple independent "
        "failure or exposure causes into a single top-level probability."
    )
    doc.add_paragraph(
        "This gives the new formula a formal mathematical grounding that the "
        "previous fixed-denominator normalisation (÷ 60.7) did not have — the "
        "60.7 constant was simply the sum of all WFVT Data Set CVSS values for this "
        "particular dataset, with no independent theoretical justification for "
        "its use as a divisor."
    )

    # ---- 9. continuous lifecycle ----
    add_heading(doc, "9. When the Score Is Recalculated", level=1)
    doc.add_paragraph(
        "The score is recalculated from scratch — over the complete current set "
        "of present findings — at each of the following points, so it always "
        "reflects the live state of the network rather than an accumulated "
        "history:"
    )
    for bullet in [
        "When an initial scan completes and configuration vulnerabilities are "
        "recorded.",
        "On every polling cycle of continuous threat detection, as active "
        "threats are newly detected or previously-active threats clear.",
        "When continuous detection is stopped, or if the field probe's "
        "heartbeat lapses.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")

    doc.add_paragraph(
        "One consequence worth highlighting to reviewers: because the score is "
        "always recomputed from the current set of present findings, it can go "
        "down as well as up — if an active threat (such as an evil-twin AP or a "
        "deauthentication attack) is cleared, its contribution is removed from the "
        "product on the very next recalculation. This is intentional: the score is "
        "designed to represent the network's risk right now, not the highest risk "
        "it has ever reached during a session."
    )

    # ---- 10. summary ----
    add_heading(doc, "10. Summary for Reviewers", level=1)
    doc.add_paragraph(
        "In summary, the methodology change preserves every element of the "
        "scoring system that was already standards-aligned — the WFVT Data Set of "
        "seven items, the use of CVSS v4.0 Base Scores as the severity input "
        "for each item, the evidence-based presence/absence determination for each "
        "finding, and the four-category LOW/MEDIUM/HIGH/CRITICAL output scale. The "
        "only element that changed is how multiple present findings are combined "
        "into one overall number: from a fixed-denominator weighted sum to a "
        "noisy-OR probabilistic combination, a change supported by NIST SP 800-30's "
        "guidance on aggregated risk, OWASP's likelihood × impact framing and "
        "single-exploit-path reasoning, and the established probability-theory / "
        "fault-tree formula for combining independent risks."
    )
    doc.add_paragraph(
        "Reviewers are invited to evaluate, in particular: (1) whether treating "
        "CVSS/10 as a per-finding compromise probability is a reasonable "
        "interpretation for this purpose, (2) whether the resulting category "
        "boundaries (e.g. a single Critical-rated finding alone reaching the "
        "CRITICAL band) match expert intuition about real-world Wi-Fi risk, and "
        "(3) whether the fixed seven-item WFVT Data Set remains an appropriate and "
        "sufficiently complete scope for the threats and vulnerabilities most "
        "relevant to the deployment environment."
    )

    doc.add_paragraph()
    add_caption(doc, "End of document.")

    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
